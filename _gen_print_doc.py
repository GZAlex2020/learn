# -*- coding: utf-8 -*-
"""将语文 / 数学打印版分别整合为 Word 文档（支持看图插图）"""
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from _print_images import demo_picture, picture_for_line

ROOT = os.path.dirname(__file__)
OUT_DIR = os.path.join(ROOT, '打印版')
DAILY_MD_MARKER = '04-每日一练'
_last_was_page_break = False

YUWEN_PARTS = [
    ('01  拼音闯关（第1～4周）', '语文/打印版/01-拼音巩固/拼音闯关练习（第1-4周）.md'),
    ('02  读题闯关（第5～8周）', '语文/打印版/03-读题训练/读题闯关练习（第5-8周）.md'),
    ('03  题型认读（选择·填空·连线·判断）', '语文/打印版/03-读题训练/题型认读-选择填空连线判断.md'),
    ('04  暑假56天每日一练', '语文/打印版/04-每日一练/暑假56天每日一练.md'),
]

MATH_PARTS = [
    ('01  口算闯关（第1～4周）', '数学/打印版/01-数感口算/口算闯关练习（第1-4周）.md'),
    ('02  读题闯关（第5～8周）', '数学/打印版/03-读题训练/读题闯关练习（第5-8周）.md'),
    ('03  题型认读（选择·填空·列式）', '数学/打印版/03-读题训练/题型认读-选择填空列式.md'),
    ('04  暑假56天每日一练', '数学/打印版/04-每日一练/暑假56天每日一练.md'),
]

# 学生版整行隐藏（含缩进的 **参考：…**）
STUDENT_HIDE_LINE_RE = re.compile(
    r'^\s*\*\*(?:答案|参考(?:列式/)?答案|参考|孩子应说|列式)[:：]'
)
ANSWER_INLINE_RE = re.compile(
    r'\*\*(?:答案|参考(?:列式/)?答案|参考)[:：][^*]*\*\*'
)
CORRECT_MARK_RE = re.compile(r'[✓✔√]\s*$|^\s*[✓✔√]\s*')
IMAGE_MD_RE = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')


def should_skip_line(line, include_answers):
    if include_answers:
        return False
    return bool(STUDENT_HIDE_LINE_RE.match(line))


def clean_student_line(line):
    line = ANSWER_INLINE_RE.sub('', line)
    line = CORRECT_MARK_RE.sub('', line)
    line = re.sub(r'\s+✓', '', line)
    line = re.sub(r'✓\s*', '', line)
    return line.rstrip()


def set_run_font(run, name='宋体', size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def mark_content_added():
    global _last_was_page_break
    _last_was_page_break = False


def add_formatted_paragraph(doc, text, style='Normal', size=11, bold=False, indent=0, space_after=3):
    if not text.strip():
        return None
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.25
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_run_font(run, bold=True, size=size)
        else:
            run = p.add_run(part)
            set_run_font(run, bold=bold, size=size)
    mark_content_added()
    return p


def add_heading(doc, text, level):
    styles = {1: ('Heading 1', 18), 2: ('Heading 2', 15), 3: ('Heading 3', 13), 4: ('Heading 4', 12)}
    style, size = styles.get(level, ('Heading 3', 12))
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(8 if level <= 2 else 4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = level >= 3
    run = p.add_run(text)
    set_run_font(run, name='黑体', size=size, bold=True)
    mark_content_added()
    return p


def add_page_break_safe(doc):
    global _last_was_page_break
    if _last_was_page_break:
        return
    if not doc.paragraphs:
        doc.add_page_break()
        _last_was_page_break = True
        return
    last = doc.paragraphs[-1]._element
    if last.tag.endswith('p'):
        brs = last.findall('.//' + qn('w:br'))
        if any(b.get(qn('w:type')) == 'page' for b in brs):
            return
    doc.add_page_break()
    _last_was_page_break = True


def reset_page_break_flag():
    global _last_was_page_break
    _last_was_page_break = False


def add_picture_if_needed(doc, text):
    img_path = picture_for_line(text)
    if not img_path or not os.path.isfile(img_path):
        if '「看图」' in text or re.search(r'「看图」是什么意思', text):
            img_path = demo_picture()
        else:
            return
    if not img_path or not os.path.isfile(img_path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Cm(10))
    mark_content_added()


def parse_table_row(line):
    line = line.strip()
    if not line.startswith('|'):
        return None
    return [c.strip() for c in line.strip('|').split('|')]


def is_separator_row(cells):
    return all(re.match(r'^:?-+:?$', c.replace(' ', '')) for c in cells if c)


def sanitize_table_rows(rows, include_answers):
    if include_answers or not rows:
        return rows
    header = rows[0]
    blank_cols = [
        i for i, cell in enumerate(header)
        if '你的答案' in cell or '孩子应说' in cell or cell.strip() in ('答案', '参考答案')
    ]
    if not blank_cols:
        return rows
    cleaned = [header]
    for row in rows[1:]:
        row = list(row)
        for ci in blank_cols:
            if ci < len(row):
                row[ci] = ''
        cleaned.append(row)
    return cleaned


def add_table(doc, rows, include_answers=True):
    rows = sanitize_table_rows(rows, include_answers)
    if len(rows) < 1:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(cols):
            text = row[ci] if ci < len(row) else ''
            cell = table.rows[ri].cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            parts = re.split(r'(\*\*[^*]+\*\*)', text)
            for part in parts:
                if not part:
                    continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    set_run_font(run, bold=True, size=10)
                else:
                    run = p.add_run(part)
                    set_run_font(run, size=10)
    mark_content_added()


def is_daily_md(md_path):
    return DAILY_MD_MARKER in md_path.replace('\\', '/')


def should_page_break(line, md_path):
    if not is_daily_md(md_path):
        return False
    if not (line.startswith('## 第 ') and '天' in line):
        return False
    m = re.search(r'第\s*(\d+)\s*天', line)
    if m and int(m.group(1)) <= 1:
        return False
    return True


def convert_md_to_doc(doc, md_path, include_answers=True):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    table_rows = []
    in_code = False

    while i < len(lines):
        raw = lines[i].rstrip('\n')
        line = raw.rstrip()

        if line.strip().startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_formatted_paragraph(doc, line, size=10)
            i += 1
            continue

        if table_rows:
            row = parse_table_row(line)
            if row:
                if not is_separator_row(row):
                    table_rows.append(row)
                i += 1
                continue
            add_table(doc, table_rows, include_answers)
            table_rows = []

        if not line.strip():
            i += 1
            continue

        if line.strip() == '---':
            i += 1
            continue

        img_m = IMAGE_MD_RE.match(line.strip())
        if img_m:
            rel = img_m.group(2)
            img_path = os.path.join(os.path.dirname(md_path), rel.replace('/', os.sep))
            if os.path.isfile(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img_path, width=Cm(10))
                mark_content_added()
            i += 1
            continue

        if should_skip_line(line, include_answers):
            i += 1
            continue

        if not include_answers:
            line = clean_student_line(line)
            if not line.strip():
                i += 1
                continue

        if line.startswith('# '):
            add_heading(doc, line[2:].strip(), 2)
            i += 1
            continue
        if line.startswith('## '):
            if should_page_break(line, md_path):
                add_page_break_safe(doc)
            add_heading(doc, line[3:].strip(), 3)
            i += 1
            continue
        if line.startswith('### '):
            add_heading(doc, line[4:].strip(), 4)
            i += 1
            continue

        if line.startswith('> '):
            add_formatted_paragraph(doc, line[2:].strip(), size=10, indent=0.5)
            i += 1
            continue

        row = parse_table_row(line)
        if row and not is_separator_row(row):
            table_rows = [row]
            i += 1
            continue

        if re.match(r'^\d+\.\s', line):
            add_formatted_paragraph(doc, line, size=11)
            i += 1
            continue

        if line.startswith('- '):
            add_formatted_paragraph(doc, '• ' + line[2:], size=11, indent=0.5)
            i += 1
            continue

        add_formatted_paragraph(doc, line, size=11)
        add_picture_if_needed(doc, line)
        i += 1

    if table_rows:
        add_table(doc, table_rows, include_answers)


def setup_document(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)


def add_cover(doc, title, subtitle, tagline):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, name='黑体', size=26, bold=True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(subtitle)
    set_run_font(run2, size=14)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(tagline)
    set_run_font(run3, size=12, color=RGBColor(80, 80, 80))
    add_page_break_safe(doc)


def build_doc(filename, parts, subject, include_answers):
    reset_page_break_flag()
    doc = Document()
    setup_document(doc)

    subj_name = '语文' if subject == 'yuwen' else '数学'
    if include_answers:
        add_cover(doc, f'一升二暑假练习册 · {subj_name}', '（家长版 · 含参考答案）', '一升二暑假复习')
        tips = [
            f'本册为{subj_name}全部闯关练习与 56 天每日一练，建议按周次顺序使用。',
            '每日一练按「第 X 天」换页，方便按天打印。',
            '读题类练习请先让孩子大声读出题目要求，再动笔或口答。',
            '含参考答案，供家长核对或批改。',
        ]
    else:
        add_cover(doc, f'一升二暑假练习册 · {subj_name}', '（学生版 · 不含答案）', '一升二暑假复习')
        tips = [
            f'本册为{subj_name}练习，不含答案，可直接打印给孩子做题。',
            '每日一练按「第 X 天」换页，方便按天打印。',
            '「看图列式」题请先看图再列算式。',
            '请用《含答案》版核对，或由家长现场批改。',
        ]

    add_heading(doc, '使用说明', 1)
    for t in tips:
        add_formatted_paragraph(doc, '• ' + t, size=11)
    add_page_break_safe(doc)

    for part_title, rel_path in parts:
        md_path = os.path.join(ROOT, rel_path.replace('/', os.sep))
        if not os.path.exists(md_path):
            print('  跳过缺失文件:', rel_path)
            continue
        add_heading(doc, part_title, 2)
        print('  写入:', rel_path)
        convert_md_to_doc(doc, md_path, include_answers=include_answers)

    out_path = os.path.join(OUT_DIR, filename)
    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    print('生成 Word 练习册（语文 / 数学分开）...')
    jobs = [
        ('一升二暑假练习册-语文（学生版）.docx', YUWEN_PARTS, 'yuwen', False),
        ('一升二暑假练习册-语文（含答案）.docx', YUWEN_PARTS, 'yuwen', True),
        ('一升二暑假练习册-数学（学生版）.docx', MATH_PARTS, 'math', False),
        ('一升二暑假练习册-数学（含答案）.docx', MATH_PARTS, 'math', True),
    ]
    for filename, parts, subject, with_ans in jobs:
        path = build_doc(filename, parts, subject, with_ans)
        print('OK', path)
    print('完成！文件位于:', OUT_DIR)


if __name__ == '__main__':
    main()
